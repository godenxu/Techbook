import os
import pptx
import io
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

pptx_path = r"C:\Users\Administrator\Documents\Antigravity\Techbook\前沿技术研究成果集-浦发银行可编辑演示文稿(已更新).pptx"
prs = pptx.Presentation(pptx_path)

print(f"Total slides in updated presentation: {len(prs.slides)}")

# Slide 15 check
slide15 = prs.slides[14]
print(f"\n--- Slide 15 ({slide15.slide_layout.name}) ---")
for s in slide15.shapes:
    if s.has_text_frame:
        p0 = s.text_frame.paragraphs[0].text.strip()
        if p0:
            print(f"  Text shape: {p0[:50]}")
    elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        print(f"  Picture shape: {s.name} (width={s.width}, height={s.height})")

# Slide 17 check
slide17 = prs.slides[16]
print(f"\n--- Slide 17 ({slide17.slide_layout.name}) ---")
for s in slide17.shapes:
    if s.has_text_frame:
        p0 = s.text_frame.paragraphs[0].text.strip()
        if p0:
            print(f"  Text shape: {p0[:50]}")
    elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        print(f"  Picture shape: {s.name} (width={s.width}, height={s.height})")

print("\nVerification completed successfully!")
