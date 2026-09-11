import os
import pptx
import io
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

pptx_path = r"C:\Users\Administrator\Documents\Antigravity\Techbook\前沿技术研究成果集-浦发银行可编辑演示文稿.pptx"

if not os.path.exists(pptx_path):
    print("ERROR: File does not exist:", pptx_path)
    sys.exit(1)

file_size = os.path.getsize(pptx_path) / (1024 * 1024)
print(f"File verified: {pptx_path} ({file_size:.2f} MB)")

prs = pptx.Presentation(pptx_path)
print(f"Total slides loaded: {len(prs.slides)}")

assert len(prs.slides) == 60, f"Expected 60 slides, got {len(prs.slides)}"

for idx, slide in enumerate(prs.slides):
    shapes_cnt = len(slide.shapes)
    tables_cnt = sum(1 for s in slide.shapes if s.has_table)
    pictures_cnt = sum(1 for s in slide.shapes if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE)
    text_previews = []
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                t = p.text.strip()
                if t and len(t) > 3:
                    text_previews.append(t)
                    if len(text_previews) >= 3:
                        break
            if len(text_previews) >= 3:
                break
    title_preview = " | ".join(text_previews[:2]) if text_previews else "(No text)"
    if len(title_preview) > 50:
        title_preview = title_preview[:50] + "..."
    print(f"Slide {idx+1:02d} | Layout: {slide.slide_layout.name:<10} | Shapes: {shapes_cnt:2d} (Tbl:{tables_cnt}, Pic:{pictures_cnt}) | Content: {title_preview}")

# Verify Slide 14 table
slide14 = prs.slides[13]
tables = [s.table for s in slide14.shapes if s.has_table]
print(f"\nSlide 14 Table Check: Found {len(tables)} tables")
if tables:
    t = tables[0]
    print(f"Table dimensions: {len(t.rows)} rows x {len(t.columns)} cols")
    print("Row 1 (Header):", [t.cell(0, c).text.strip() for c in range(min(5, len(t.columns)))])
    print("Row 2 (T01):", [t.cell(1, c).text.strip() for c in range(min(5, len(t.columns)))])
    print("Row 37 (T36):", [t.cell(36, c).text.strip() for c in range(min(5, len(t.columns)))])

# Verify Slide 19 (T01)
slide19 = prs.slides[18]
print(f"\nSlide 19 (T01) Check: Shapes count = {len(slide19.shapes)}")
for s_idx, s in enumerate(slide19.shapes):
    if s.has_text_frame:
        p0 = s.text_frame.paragraphs[0].text.strip()
        if p0:
            print(f"  Shape {s_idx} text: {p0[:40]}")
    elif s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        print(f"  Shape {s_idx} PICTURE: {s.name} (width={s.width}, height={s.height})")

print("\nALL 60 SLIDES VERIFIED SUCCESSFULLY!")
